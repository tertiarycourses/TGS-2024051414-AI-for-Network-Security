#!/usr/bin/env python3
"""Generate the AI for Network Security (TGS-2024051414) Learner Guide as BOTH a Markdown mirror
(LG-*.md at repo root) and a DOCX (courseware/LG-*.docx) from one source, so they never diverge.

House format: cover page, Document Version Control Record, auto TOC, Arial 11pt
body, one section per lab (Objective · Goal · What you'll build · Step-by-step
with commands · Test it), plus setup, exam-prep and glossary. All content is
driven by course_data + the domain data files, keeping the LG 100% aligned with
the slide deck, Lesson Plan and labs.
"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches as DocxInches
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

# ---------------- block DSL (single content stream → MD + DOCX) ----------------
B=[]
def h1(t): B.append(("h1",t))
def h2(t): B.append(("h2",t))
def h3(t): B.append(("h3",t))
def p(t):  B.append(("p",t))
def bullets(xs): B.append(("bullets",xs))
def steps(xs): B.append(("steps",xs))
def code(t): B.append(("code",t))
def note(t): B.append(("note",t))
def rule(): B.append(("rule",))
def figure(path, caption): B.append(("figure", path, caption))

# ---------------- content ----------------
h1("Introduction")
p(f"This Learner Guide accompanies the WSQ course {C.TITLE} ({C.COURSE_CODE}), conducted by {C.ORG}. "
  "It provides the full step-by-step instructions for all 12 hands-on labs, organised by the four "
  f"course topics and mapped to the Skills Framework TSC {C.TSC_TITLE} ({C.TSC_CODE}). Every lab maps "
  "to a learning outcome and to the specific knowledge (K) and ability (A) statements that the "
  "assessment measures.")
p("The course teaches a single, practical proposition: attacks now move at machine speed and at a "
  "volume no analyst can read, so the defence has to be assisted by machines too. You will use AI to "
  "analyse network configurations, baseline traffic, detect anomalies, prioritise vulnerabilities, "
  "score risk, design and implement controls, triage alerts and document incidents. You will also "
  "learn where AI must not be trusted — and how to attack and then harden an AI model yourself.")
p("Use this guide alongside the course slides and the lab folders in labs/. The slides carry the "
  "concepts and the shape of each lab; this guide carries the detailed procedure. Every scanning or "
  "testing technique in this course must only be run against systems you own or have written "
  "authorisation to test. Unauthorised access to a computer is a criminal offence under the Singapore "
  "Computer Misuse Act and its equivalents elsewhere.")

h1("Course Learning Outcomes")
bullets(C.LEARNING_OUTCOMES)

h1("Underpinning Knowledge and Abilities Assessed")
h3("Knowledge (K) — assessed by the Written Assessment (WA-SAQ)")
bullets([f"{k}: {v}" for k, v in C.KNOWLEDGE])
h3("Abilities (A) — assessed by the Case Study (CS)")
bullets([f"{a}: {v}" for a, v in C.ABILITIES])
h3("How the topics map to the outcomes")
bullets([
 f"Topic {t['code']} — {t['title']}: {', '.join(t['los'])} · knowledge {', '.join(t['ks'])} · "
 f"abilities {', '.join(t['abil'])} · labs {', '.join(str(x['num']) for x in ACT if x['topic']==t['num'])}"
 for t in C.TOPICS])

h1("Before You Start — Environment Setup")
h3("What you need")
bullets([f"{n} — {d}" for n, d in C.LAB_TOOLS])
h3("Install and verify the toolchain")
p("Install the packet-capture, scanning and Python tooling before Lab 1. The commands below cover "
  "macOS (Homebrew), Debian/Ubuntu Linux and Windows; use whichever matches your laptop. Verify each "
  "tool reports a version before you continue — a lab that fails at step 1 usually fails because a "
  "tool was never actually installed.")
code("# macOS (Homebrew)\n"
     "brew install nmap wireshark python@3.12\n\n"
     "# Debian / Ubuntu Linux\n"
     "sudo apt update && sudo apt install -y nmap wireshark tshark python3 python3-pip\n\n"
     "# Windows (winget, from an elevated PowerShell)\n"
     "winget install Insecure.Nmap WiresharkFoundation.Wireshark Python.Python.3.12\n\n"
     "# Verify — every one of these must print a version\n"
     "nmap --version\n"
     "tshark --version\n"
     "python3 --version")
h3("Set up the Python environment for the AI labs")
p("Labs 2, 3, 9 and 10 use pandas and scikit-learn. Create an isolated virtual environment so the "
  "course packages never collide with anything else on your machine.")
code("# create and activate a virtual environment\n"
     "python3 -m venv ~/ainetsec-venv\n"
     "source ~/ainetsec-venv/bin/activate      # macOS / Linux\n"
     "# .\\ainetsec-venv\\Scripts\\Activate.ps1   # Windows PowerShell\n\n"
     "# install the packages used by the labs\n"
     "pip install --upgrade pip\n"
     "pip install pandas scikit-learn matplotlib jupyter\n\n"
     "# verify\n"
     "python3 -c \"import pandas, sklearn; print('pandas', pandas.__version__, '| scikit-learn', sklearn.__version__)\"")
h3("Set up your AI assistant")
p("Several labs use a generative AI assistant (Claude, ChatGPT or Copilot) as an analyst copilot. You "
  "do not need a paid plan — the free tier of any of them is sufficient. Two rules apply throughout "
  "the course and are themselves assessable good practice:")
bullets([
 "Never paste real production data, credentials, customer information or unredacted logs into a public AI assistant. Sanitise first — replace real addresses, hostnames and identifiers with lab equivalents.",
 "Always verify what the assistant tells you against an authoritative source. Every lab that uses AI has a verification step for exactly this reason; a confident answer is not a correct answer.",
])
h3("Authorisation — read this before Lab 1")
p("Labs 1, 2, 3 and 8 run active scanning and packet capture. Run them only against the lab network "
  "your trainer provides, or against a network you own. If you intend to repeat these labs at work, "
  "get written authorisation from the system owner first, defining which addresses are in scope and "
  "during which window. That written authorisation is what separates a security assessment from an "
  "offence.")
h3("Conventions used in every lab")
bullets([
 "Commands are shown for a terminal (macOS/Linux) unless a Windows form is given; a leading sudo indicates administrator rights are required.",
 "Placeholders such as 192.168.56.0/24 and 192.168.56.101 are the lab network and lab host — replace them with the addresses your trainer gives you.",
 "Each lab produces a named deliverable. Keep every deliverable — later labs build directly on earlier ones, and the Case Study assessment rehearses the same chain.",
 "Where a lab asks you to compare your own judgement with the AI's output, record both. The comparison, not the AI's answer, is the learning.",
])

h1("How the Labs Build on Each Other")
p("The twelve labs are one continuous engagement, not twelve unrelated exercises. Each deliverable "
  "feeds the next lab, and the whole chain is what the Case Study assessment asks you to reproduce in "
  "miniature. Losing an early deliverable will block a later lab.")
bullets([
 "Lab 1 produces the network configuration inventory — the raw material for Labs 3 and 4.",
 "Lab 2 produces the traffic baseline and the anomaly model — reused and then attacked in Lab 9.",
 "Lab 3 turns the inventory into a prioritised vulnerability worklist — the input to the risk register.",
 "Lab 4 turns the inventory into a classified asset register — the other input to the risk register.",
 "Lab 5 joins Labs 3 and 4 into a scored risk register with inherent and residual risk.",
 "Lab 6 turns the register into a treated risk management plan mapped to the NIST CSF.",
 "Lab 7 turns the plan into a three-phase segmented security design with a zone access matrix.",
 "Lab 8 turns the design into an implemented default-deny ruleset with tuned detection rules.",
 "Lab 9 attacks the AI you have been relying on, then hardens it against those attacks.",
 "Lab 10 specifies the logging that makes response possible, and triages the alert queue with AI.",
 "Lab 11 documents the top incident from Lab 10 to an evidentiary standard.",
 "Lab 12 implements and tests the response, then feeds the control changes back into the Lab 6 plan.",
])

# ---------------- per-topic, per-lab ----------------
TOPICS_BY_NUM = {t["num"]: t for t in C.TOPICS}
for t in C.TOPICS:
    h1(f"Topic {t['code']} — {t['title']}  ({t['weighting']})")
    p(t["subtitle"])
    p(f"Learning outcome: {', '.join(t['los'])}.  "
      f"Knowledge assessed: {', '.join(t['ks'])}.  "
      f"Abilities assessed: {', '.join(t['abil'])}.")
    h3("Key concepts")
    bullets(t["concepts"])
    for a in [x for x in ACT if x["topic"] == t["num"]]:
        h2(f"Lab {a['num']} — {a['title']}")
        p(f"Maps to: {a['objective']}")
        p(f"Goal: {a['desc']}")
        _fig = os.path.join(REPO, "courseware", "assets", "labs", f"lab-{a['num']:02d}-process.png")
        if os.path.exists(_fig):
            figure(_fig, f"Figure {a['num']} — the Lab {a['num']} workflow: "
                         + " → ".join(a.get("stages") or []) + ".")
        h3("What you'll produce")
        p(a["build"] + f"   (Tools: {a['services']}.)")
        h3("Step-by-step")
        steps([(instr, cmd) for instr, cmd in a["steps"]])
        h3("Test it — how you know the lab is done")
        p(a["test"])
        if a.get("troubleshoot"):
            h3("If it doesn't work")
            p(a["troubleshoot"])
        note(f"Only run these techniques against systems you own or are authorised in writing to test. "
             f"Keep the deliverable from this lab — later labs build on it.")
        rule()

h1("Assessment Preparation")
h3("What the assessment consists of")
bullets([
 C.ASSESSMENT["written"],
 C.ASSESSMENT["practical"],
 "Both instruments are open book: you may use these course slides, this Learner Guide and approved materials only.",
 "You are assessed as Competent (C) or Not Yet Competent (NYC) on each instrument. An NYC in any K or A results in an NYC for the unit.",
 C.ASSESSMENT["note"],
])
h3("Preparing for the Written Assessment (K1–K4)")
bullets([
 "K1 — be able to describe an ICT network's components and configuration, and explain what a given configuration weakness exposes. Revise Topic 1 and your Lab 1 inventory.",
 "K2 — be able to name and describe the main classes of network attack and the vulnerabilities they exploit, including why polymorphic and Living-off-the-Land techniques defeat signature detection. Revise Topic 1 and Topic 2.",
 "K3 — be able to describe the types and techniques of network security measures across the physical, technical and administrative categories, and the preventive/detective/corrective distinction. Revise Topic 3.",
 "K4 — be able to explain how a risk management plan is built and how incident response procedures are structured, including the NIST SP 800-61 lifecycle. Revise Topics 2 and 4.",
 "Practise writing short, complete answers: state the point, then justify it with a reason or an example. A one-word answer rarely evidences the knowledge.",
])
h3("Preparing for the Case Study (A1–A9)")
p("The Case Study presents a realistic organisational scenario and asks you to work through the same "
  "chain you practised in the labs. The most reliable preparation is to redo the chain end to end on "
  "your own lab network, from memory, and check that each deliverable still comes out complete.")
bullets([
 "A1, A2 — identify the threats in the scenario and analyse the security risks, using evidence from the material given rather than generic assertions.",
 "A3, A4 — determine which organisational assets require protection and create a risk management plan with treatments, owners and dates.",
 "A5, A6, A7 — define the planning, building and management phases, develop the security measures for the network components, and design the auditing and incident response procedures.",
 "A8, A9 — document the security incident described, and state the incident response configurations you would implement.",
 "Answer in the scenario's own terms. A generic answer that could apply to any organisation does not evidence the ability.",
])
h3("Practice exam")
bullets([
 f"Attempt the Tertiary Infotech practice exam before the assessment: {C.EXAM_URL}",
 "Sit it under exam conditions — timed and closed book — so the score tells you something real.",
 "Trace every wrong answer back to the slide and the Learner Guide section that covers it, then re-attempt.",
])

h1("Reference Sources")
p("The AI content in this course is grounded in the following authoritative and industry sources. "
  "They are worth reading in full after the course.")
bullets([
 "NIST Cybersecurity Framework 2.0 — Govern, Identify, Protect, Detect, Respond, Recover: https://www.nist.gov/cyberframework",
 "NIST AI 100-2e2025, Adversarial Machine Learning: A Taxonomy and Terminology of Attacks and Mitigations: https://csrc.nist.gov/pubs/ai/100/2/e2025/final",
 "NSA / CISA / FBI joint guidance, Deploying AI Systems Securely (April 2024): https://www.cisa.gov/news-events/alerts/2024/04/15/joint-guidance-deploying-ai-systems-securely",
 "CISA Artificial Intelligence programme and AI data security guidance: https://www.cisa.gov/ai",
 "MITRE ATLAS — adversary tactics and techniques against AI systems: https://atlas.mitre.org/",
 "NIST SP 800-61, Computer Security Incident Handling Guide: https://csrc.nist.gov/pubs/sp/800/61/r2/final",
 "CISA Known Exploited Vulnerabilities catalogue: https://www.cisa.gov/known-exploited-vulnerabilities-catalog",
 "NIST National Vulnerability Database: https://nvd.nist.gov/",
 "Check Point, AI in Network Security: https://www.checkpoint.com/cyber-hub/cyber-security/what-is-ai-security/ai-in-network-security/",
 "Fortinet, Artificial Intelligence in Cybersecurity: https://www.fortinet.com/resources/cyberglossary/artificial-intelligence-in-cybersecurity",
 "PDPC Singapore — data breach notification obligation under the PDPA: https://www.pdpc.gov.sg/",
])

h1("Glossary")
gl = [
 ("Anomaly detection", "Finding events that deviate from a learned baseline of normal behaviour, rather than matching a known bad pattern."),
 ("Pattern of life", "The learned profile of an entity's normal behaviour — its usual peers, ports, volumes and hours — against which deviations are scored."),
 ("NDR", "Network Detection and Response — a platform applying behavioural analytics to network traffic across the attack lifecycle."),
 ("SIEM", "Security Information and Event Management — collects and correlates log events from across the estate into incidents."),
 ("SOAR", "Security Orchestration, Automation and Response — executes the response playbook automatically once an incident is raised."),
 ("UEBA", "User and Entity Behaviour Analytics — profiles each user and device against its own history to detect deviation."),
 ("CIA triad", "Confidentiality, Integrity and Availability — the three goals every security control ultimately serves."),
 ("Defence in depth", "Layering physical, technical and administrative controls so no single failure becomes a breach."),
 ("Zero trust", "An architecture that never grants access on the basis of network location alone — every request is verified."),
 ("Segmentation", "Dividing a network into zones with controlled boundaries, so a compromise cannot move freely."),
 ("Default deny", "A policy that blocks all traffic except what is explicitly permitted — the rule that makes segmentation real."),
 ("CVSS", "Common Vulnerability Scoring System — a 0–10 severity score for a vulnerability."),
 ("KEV", "CISA's Known Exploited Vulnerabilities catalogue — vulnerabilities confirmed to be exploited in the wild."),
 ("Inherent vs residual risk", "Risk before any control is applied, versus the risk that remains after the controls are in place."),
 ("Risk treatment", "The decision taken on a risk: mitigate, transfer, avoid or accept."),
 ("Evasion attack", "Crafting input so a deployed model misclassifies it — for example slowing a scan until it falls inside the baseline."),
 ("Poisoning attack", "Corrupting the training data or the model so that its future judgements are degraded in the attacker's favour."),
 ("Isolation Forest", "An unsupervised algorithm that isolates outliers, used in this course for network anomaly detection."),
 ("False positive / false negative", "A benign event wrongly flagged / a malicious event wrongly missed. Tuning trades one against the other."),
 ("Human in the loop", "A required human decision before an automated action is taken, used where the action's blast radius is large."),
 ("Blast radius", "The extent of what an action affects — the test for whether a response step may be automated."),
 ("Chain of custody", "The documented handling of evidence that keeps it defensible for audit, regulator or court."),
 ("MTTD / MTTR", "Mean Time To Detect / Mean Time To Respond — the two headline measures of a security operation."),
]
B.append(("dl", gl))

# ---------------- render Markdown ----------------
def _anchor(txt):
    return "".join(ch.lower() if ch.isalnum() else ("-" if ch in " -" else "") for ch in txt)

def render_md():
    out=[f"# {C.TITLE} — Learner Guide",""]
    out.append(f"**WSQ Course Code:** {C.COURSE_CODE}  |  **Conducted by:** {C.ORG} ({C.UEN.replace('UEN: ','UEN ')})  |  **Version {C.VERSION} · {C.VERSION_DATE}**")
    out.append("")
    # TOC (h1 + h2)
    out.append("## Contents"); out.append("")
    for kind,*rest in B:
        if kind=="h1": out.append(f"- [{rest[0]}](#{_anchor(rest[0])})")
        elif kind=="h2": out.append(f"  - [{rest[0]}](#{_anchor(rest[0])})")
    out.append("")
    for kind,*rest in B:
        if kind=="h1": out+=["",f"## {rest[0]}",""]
        elif kind=="h2": out+=["",f"### {rest[0]}",""]
        elif kind=="h3": out+=[f"**{rest[0]}**",""]
        elif kind=="p": out+=[rest[0],""]
        elif kind=="bullets": out+=[f"- {x}" for x in rest[0]]+[""]
        elif kind=="steps":
            for i,(instr,cmd) in enumerate(rest[0],1):
                out.append(f"{i}. {instr}")
                if cmd: out+=["",f"   ```bash",f"   {cmd}","   ```",""]
            out.append("")
        elif kind=="code": out+=["```bash",rest[0],"```",""]
        elif kind=="note": out+=[f"> **Note:** {rest[0]}",""]
        elif kind=="rule": out+=["---",""]
        elif kind=="figure":
            _rel=os.path.relpath(rest[0], REPO)
            out+=[f"![{rest[1]}]({_rel})","",f"*{rest[1]}*",""]
        elif kind=="dl":
            for term,defn in rest[0]: out.append(f"- **{term}** — {defn}")
            out.append("")
    return "\n".join(out)

MD_OUT=os.path.join(REPO,f"LG-{C.SHORT_TITLE}.md")
with open(MD_OUT,"w") as f: f.write(render_md())
print("Saved",MD_OUT)

# ---------------- render DOCX ----------------
BRAND=RGBColor(0x1F,0x6F,0xEB); DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)
INKCODE=RGBColor(0x0B,0x30,0x60)
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)
prodoc.add_cover_page(doc,"LEARNER GUIDE",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("1.0",C.VERSION_DATE,
  "Initial release of the AI for Network Security Learner Guide — 4 topics aligned to LO1-LO4 and TSC "
  +C.TSC_CODE+", with full step-by-step procedures for all 12 AI-assisted hands-on labs, the K1-K4 / "
  "A1-A9 mapping, assessment preparation for the WA-SAQ and Case Study, and the reference source list.",
  C.TRAINER),
 (C.VERSION.lstrip("v"),C.VERSION_DATE,
  "Quality-audit revision — regenerated in step with the v1.1 slide deck and lesson plan so the "
  "lab procedures, slide references and assessment timings stay identical across all artifacts.",
  C.TRAINER),
])
prodoc.add_toc(doc)

def code_para(text):
    for line in text.split("\n"):
        para=doc.add_paragraph(); prodoc._shade_para(para) if hasattr(prodoc,"_shade_para") else None
        r=para.add_run(line); r.font.name="Consolas"; r.font.size=Pt(9.5); r.font.color.rgb=INKCODE

for kind,*rest in B:
    if kind=="h1": doc.add_heading(rest[0],level=1)
    elif kind=="h2": doc.add_heading(rest[0],level=2)
    elif kind=="h3":
        para=doc.add_paragraph(); r=para.add_run(rest[0]); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=BRAND
    elif kind=="p": doc.add_paragraph(rest[0])
    elif kind=="bullets":
        for x in rest[0]: doc.add_paragraph(x,style="List Bullet")
    elif kind=="steps":
        for i,(instr,cmd) in enumerate(rest[0],1):
            para=doc.add_paragraph(style="List Number"); para.add_run(instr)
            if cmd: code_para(cmd)
    elif kind=="code": code_para(rest[0])
    elif kind=="note":
        para=doc.add_paragraph(); r=para.add_run("Note: "); r.bold=True; r.font.color.rgb=BRAND
        para.add_run(rest[0]).font.size=Pt(10)
    elif kind=="figure":
        try:
            doc.add_picture(rest[0], width=DocxInches(6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = cap.add_run(rest[1]); cr.italic = True; cr.font.size = Pt(9)
            cr.font.color.rgb = GREY
        except Exception as _e:
            print("  ! figure skipped:", rest[0], _e)
    elif kind=="rule": doc.add_paragraph("")
    elif kind=="dl":
        for term,defn in rest[0]:
            para=doc.add_paragraph(style="List Bullet")
            r=para.add_run(term+" — "); r.bold=True; para.add_run(defn)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
DOCX_OUT=os.path.join(REPO,"courseware",f"LG-{C.SHORT_TITLE}.docx")
doc.save(DOCX_OUT)
print("Saved",DOCX_OUT)
